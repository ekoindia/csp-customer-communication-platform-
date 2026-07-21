import os
from typing import List, Dict


def detect_format(file_path: str) -> str:
    ext = os.path.splitext(file_path)[1].lower()
    if ext in (".xlsx", ".xls"):
        return "excel"
    if ext == ".csv":
        return "csv"
    if ext == ".pdf":
        return "pdf"
    if ext in (".jpg", ".jpeg", ".png", ".tiff", ".bmp"):
        return "image"
    raise ValueError(f"Unsupported file type: {ext}")


def parse(file_path: str) -> List[Dict]:
    fmt = detect_format(file_path)
    if fmt == "excel":
        return _parse_excel(file_path)
    if fmt == "csv":
        return _parse_csv(file_path)
    if fmt == "pdf":
        return _parse_pdf(file_path)
    if fmt == "image":
        return _parse_image(file_path)


def _parse_excel(file_path: str) -> List[Dict]:
    import openpyxl
    # close the workbook before returning: on Windows an un-closed workbook keeps
    # a handle on the file, which then blocks the DPDP cleanup from deleting the
    # raw upload afterwards (leaving customer data on disk — critical when the
    # file is a decrypted .cspx). read_only + close() releases it immediately.
    wb = openpyxl.load_workbook(file_path, data_only=True, read_only=True)
    try:
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        return []
    headers = _normalise_headers(rows[0])
    return [dict(zip(headers, row)) for row in rows[1:] if any(cell is not None for cell in row)]


def _parse_csv(file_path: str) -> List[Dict]:
    import csv
    with open(file_path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        return [dict(row) for row in reader]


def _parse_pdf(file_path: str) -> List[Dict]:
    import pdfplumber
    rows = []
    with pdfplumber.open(file_path) as pdf:
        for page in pdf.pages:
            table = page.extract_table()
            if table:
                headers = _normalise_headers(table[0])
                for row in table[1:]:
                    if any(cell for cell in row):
                        rows.append(dict(zip(headers, row)))
    if rows:
        return rows
    return _parse_scanned_pdf(file_path)


def _parse_scanned_pdf(file_path: str) -> List[Dict]:
    import gc
    import pypdfium2 as pdfium
    from core.ocr_table import extract_rows_from_pil
    from core import hardware

    # Adaptive DPI: lower on a 4 GB box so a full-page render doesn't spike RAM
    # into swap (see hardware.render_dpi). Configurable via config.OCR_RENDER_DPI.
    scale = hardware.render_dpi() / 72
    all_rows = []
    pdf = pdfium.PdfDocument(file_path)
    try:
        for page in pdf:
            image = page.render(scale=scale).to_pil()
            try:
                all_rows.extend(extract_rows_from_pil(image))
            finally:
                # Free this page's large image BEFORE rendering the next one, so
                # only ONE page ever sits in RAM at a time (critical on 4 GB).
                try:
                    image.close()
                except Exception:
                    pass
                del image
                gc.collect()
    finally:
        pdf.close()
    return all_rows


def _parse_image(file_path: str) -> List[Dict]:
    from PIL import Image
    from core.ocr_table import extract_rows_from_pil
    return extract_rows_from_pil(Image.open(file_path))


def _normalise_headers(header_row) -> List[str]:
    result = []
    for i, h in enumerate(header_row):
        if h is None:
            result.append(f"col_{i}")
        else:
            result.append(str(h).strip().lower().replace(" ", "_"))
    return result
